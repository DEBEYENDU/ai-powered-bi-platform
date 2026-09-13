"""Multi-format export engine for AI-generated reports.

Exports report content to PDF, DOCX, PPTX, XLSX, CSV, JSON, Markdown, HTML.
Uses optional dependencies (reportlab, python-docx, openpyxl, python-pptx)
with graceful fallback to HTML when unavailable.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def _build_html(report: dict[str, Any]) -> str:
    """Build a self-contained HTML report."""
    title = report.get("title", "Report")
    summary = report.get("executive_summary", "")
    sections = report.get("sections", [])
    kpis = report.get("kpis", [])
    insights = report.get("insights", [])
    risks = report.get("risks", [])
    recs = report.get("recommendations", [])

    kpi_html = ""
    if kpis:
        rows = "".join(
            f"<tr><td>{k.get('name', '')}</td><td>{k.get('value', '')}</td>"
            f"<td>{k.get('unit', '')}</td><td>{k.get('change_pct', 0):+.1f}%</td>"
            f"<td>{k.get('trend', '')}</td></tr>"
            for k in kpis
        )
        kpi_html = f"<h2>Key Performance Indicators</h2><table><tr><th>Metric</th><th>Value</th><th>Unit</th><th>Change</th><th>Trend</th></tr>{rows}</table>"

    sections_html = ""
    for s in sections:
        content = s.get("content", "")
        sections_html += f"<h2>{s.get('title', 'Section')}</h2><div>{content}</div>"

    insights_html = ""
    if insights:
        items = "".join(
            f"<li><strong>{i.get('title', '')}</strong>: {i.get('description', '')}</li>"
            for i in insights
        )
        insights_html = f"<h2>Key Insights</h2><ul>{items}</ul>"

    risks_html = ""
    if risks:
        items = "".join(
            f"<li><strong>{r.get('title', '')}</strong> [{r.get('severity', '')}]: {r.get('description', '')}</li>"
            for r in risks
        )
        risks_html = f"<h2>Risks</h2><ul>{items}</ul>"

    recs_html = ""
    if recs:
        items = "".join(
            f"<li><strong>{r.get('title', '')}</strong> [{r.get('priority', '')}]: {r.get('description', '')}</li>"
            for r in recs
        )
        recs_html = f"<h2>Recommendations</h2><ul>{items}</ul>"

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{title}</title>
<style>
body{{font-family:Arial,sans-serif;margin:40px;color:#333}}
h1{{color:#1a237e;border-bottom:2px solid #1a237e;padding-bottom:8px}}
h2{{color:#283593;margin-top:24px}}
table{{border-collapse:collapse;width:100%;margin:12px 0}}
th,td{{border:1px solid #ddd;padding:8px;text-align:left}}
th{{background:#f5f5f5;font-weight:600}}
ul{{line-height:1.8}}
.summary{{background:#f8f9fa;padding:16px;border-left:4px solid #1a237e;margin:16px 0}}
</style></head><body>
<h1>{title}</h1>
<div class="summary"><h2>Executive Summary</h2><p>{summary}</p></div>
{kpi_html}{sections_html}{insights_html}{risks_html}{recs_html}
</body></html>"""


def _build_markdown(report: dict[str, Any]) -> str:
    """Build a Markdown report."""
    lines: list[str] = []
    lines.append(f"# {report.get('title', 'Report')}\n")
    lines.append(f"## Executive Summary\n\n{report.get('executive_summary', '')}\n")

    kpis = report.get("kpis", [])
    if kpis:
        lines.append("## Key Performance Indicators\n")
        lines.append("| Metric | Value | Unit | Change | Trend |")
        lines.append("|--------|-------|------|--------|-------|")
        for k in kpis:
            lines.append(
                f"| {k.get('name', '')} | {k.get('value', '')} | {k.get('unit', '')} | {k.get('change_pct', 0):+.1f}% | {k.get('trend', '')} |"
            )
        lines.append("")

    for s in report.get("sections", []):
        lines.append(f"## {s.get('title', 'Section')}\n\n{s.get('content', '')}\n")

    insights = report.get("insights", [])
    if insights:
        lines.append("## Key Insights\n")
        for i in insights:
            lines.append(f"- **{i.get('title', '')}**: {i.get('description', '')}")
        lines.append("")

    risks = report.get("risks", [])
    if risks:
        lines.append("## Risks\n")
        for r in risks:
            lines.append(
                f"- **{r.get('title', '')}** [{r.get('severity', '')}]: {r.get('description', '')}"
            )
        lines.append("")

    recs = report.get("recommendations", [])
    if recs:
        lines.append("## Recommendations\n")
        for r in recs:
            lines.append(
                f"- **[{r.get('priority', '').upper()}]** {r.get('title', '')}: {r.get('description', '')}"
            )
        lines.append("")

    return "\n".join(lines)


def export_html(report: dict[str, Any], path: Path) -> dict[str, Any]:
    path.write_text(_build_html(report), encoding="utf-8")
    return {"format": "html", "file_size": path.stat().st_size}


def export_markdown(report: dict[str, Any], path: Path) -> dict[str, Any]:
    path.write_text(_build_markdown(report), encoding="utf-8")
    return {"format": "markdown", "file_size": path.stat().st_size}


def export_json(report: dict[str, Any], path: Path) -> dict[str, Any]:
    data = json.loads(json.dumps(report, default=str))
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return {"format": "json", "file_size": path.stat().st_size}


def export_csv(report: dict[str, Any], path: Path) -> dict[str, Any]:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Type", "Title", "Content", "Priority", "Severity", "Confidence"])
        for i in report.get("insights", []):
            writer.writerow(
                [
                    "insight",
                    i.get("title", ""),
                    i.get("description", ""),
                    "",
                    "",
                    i.get("confidence", ""),
                ]
            )
        for r in report.get("risks", []):
            writer.writerow(
                [
                    "risk",
                    r.get("title", ""),
                    r.get("description", ""),
                    "",
                    r.get("severity", ""),
                    "",
                ]
            )
        for r in report.get("recommendations", []):
            writer.writerow(
                [
                    "recommendation",
                    r.get("title", ""),
                    r.get("description", ""),
                    r.get("priority", ""),
                    "",
                    r.get("confidence", ""),
                ]
            )
        for k in report.get("kpis", []):
            writer.writerow(["kpi", k.get("name", ""), str(k.get("value", "")), "", "", ""])
    return {"format": "csv", "file_size": path.stat().st_size}


def export_pdf(report: dict[str, Any], path: Path) -> dict[str, Any]:
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # type: ignore

        doc = SimpleDocTemplate(str(path), pagesize=A4)
        styles = getSampleStyleSheet()
        story: list[Any] = []

        story.append(Paragraph(report.get("title", "Report"), styles["Title"]))
        story.append(Spacer(1, 12))
        story.append(Paragraph("Executive Summary", styles["Heading2"]))
        story.append(Paragraph(report.get("executive_summary", ""), styles["BodyText"]))
        story.append(Spacer(1, 12))

        for k in report.get("kpis", []):
            story.append(
                Paragraph(
                    f"{k.get('name', '')}: {k.get('value', '')} {k.get('unit', '')} ({k.get('change_pct', 0):+.1f}%)",
                    styles["BodyText"],
                )
            )

        story.append(Spacer(1, 12))
        for s in report.get("sections", []):
            story.append(Paragraph(s.get("title", ""), styles["Heading2"]))
            story.append(Paragraph(s.get("content", ""), styles["BodyText"]))
            story.append(Spacer(1, 8))

        recs = report.get("recommendations", [])
        if recs:
            story.append(Paragraph("Recommendations", styles["Heading2"]))
            for r in recs:
                story.append(
                    Paragraph(
                        f"[{r.get('priority', '').upper()}] {r.get('title', '')}: {r.get('description', '')}",
                        styles["BodyText"],
                    )
                )

        doc.build(story)
        return {"format": "pdf", "file_size": path.stat().st_size}
    except Exception as exc:  # noqa: BLE001
        html_path = path.with_suffix(".html")
        html_path.write_text(_build_html(report), encoding="utf-8")
        return {"format": "pdf", "degraded": True, "fallback": str(html_path), "reason": str(exc)}


def export_docx(report: dict[str, Any], path: Path) -> dict[str, Any]:
    try:
        from docx import Document  # type: ignore

        doc = Document()
        doc.add_heading(report.get("title", "Report"), level=1)
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(report.get("executive_summary", ""))

        for k in report.get("kpis", []):
            doc.add_paragraph(
                f"{k.get('name', '')}: {k.get('value', '')} {k.get('unit', '')} ({k.get('change_pct', 0):+.1f}%)"
            )

        for s in report.get("sections", []):
            doc.add_heading(s.get("title", ""), level=2)
            doc.add_paragraph(s.get("content", ""))

        recs = report.get("recommendations", [])
        if recs:
            doc.add_heading("Recommendations", level=2)
            for r in recs:
                doc.add_paragraph(
                    f"[{r.get('priority', '').upper()}] {r.get('title', '')}: {r.get('description', '')}"
                )

        doc.save(str(path))
        return {"format": "docx", "file_size": path.stat().st_size}
    except Exception as exc:  # noqa: BLE001
        html_path = path.with_suffix(".html")
        html_path.write_text(_build_html(report), encoding="utf-8")
        return {"format": "docx", "degraded": True, "fallback": str(html_path), "reason": str(exc)}


def export_xlsx(report: dict[str, Any], path: Path) -> dict[str, Any]:
    try:
        from openpyxl import Workbook  # type: ignore

        wb = Workbook()
        ws = wb.active
        ws.title = "KPIs"
        ws.append(["Metric", "Value", "Unit", "Change %", "Trend"])
        for k in report.get("kpis", []):
            ws.append(
                [
                    k.get("name", ""),
                    k.get("value", ""),
                    k.get("unit", ""),
                    k.get("change_pct", 0),
                    k.get("trend", ""),
                ]
            )

        ws2 = wb.create_sheet("Insights")
        ws2.append(["Title", "Description", "Type", "Confidence"])
        for i in report.get("insights", []):
            ws2.append(
                [
                    i.get("title", ""),
                    i.get("description", ""),
                    i.get("insight_type", ""),
                    i.get("confidence", ""),
                ]
            )

        ws3 = wb.create_sheet("Recommendations")
        ws3.append(["Title", "Description", "Priority", "Category", "Impact"])
        for r in report.get("recommendations", []):
            ws3.append(
                [
                    r.get("title", ""),
                    r.get("description", ""),
                    r.get("priority", ""),
                    r.get("category", ""),
                    r.get("expected_impact", ""),
                ]
            )

        wb.save(str(path))
        return {"format": "xlsx", "file_size": path.stat().st_size}
    except Exception as exc:  # noqa: BLE001
        html_path = path.with_suffix(".html")
        html_path.write_text(_build_html(report), encoding="utf-8")
        return {"format": "xlsx", "degraded": True, "fallback": str(html_path), "reason": str(exc)}


def export_pptx(report: dict[str, Any], path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation  # type: ignore

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = report.get("title", "Report")[:100]

        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Executive Summary"
        slide.placeholders[1].text = report.get("executive_summary", "")[:1000]

        for k in report.get("kpis", [])[:8]:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = k.get("name", "")[:100]
            slide.placeholders[
                1
            ].text = f"Value: {k.get('value', '')} {k.get('unit', '')}\nChange: {k.get('change_pct', 0):+.1f}%\nTrend: {k.get('trend', '')}"

        for s in report.get("sections", [])[:10]:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = s.get("title", "")[:100]
            slide.placeholders[1].text = s.get("content", "")[:1000]

        prs.save(str(path))
        return {"format": "pptx", "file_size": path.stat().st_size}
    except Exception as exc:  # noqa: BLE001
        html_path = path.with_suffix(".html")
        html_path.write_text(_build_html(report), encoding="utf-8")
        return {"format": "pptx", "degraded": True, "fallback": str(html_path), "reason": str(exc)}


EXPORTERS = {
    "pdf": export_pdf,
    "docx": export_docx,
    "pptx": export_pptx,
    "xlsx": export_xlsx,
    "csv": export_csv,
    "json": export_json,
    "html": export_html,
    "markdown": export_markdown,
}


def export_report(report: dict[str, Any], fmt: str, path: Path) -> dict[str, Any]:
    """Export a report in the specified format. Returns metadata dict."""
    fmt = fmt.lower()
    exporter = EXPORTERS.get(fmt)
    if not exporter:
        raise ValueError(f"Unsupported format '{fmt}'. Supported: {list(EXPORTERS.keys())}")
    path.parent.mkdir(parents=True, exist_ok=True)
    return exporter(report, path)

from __future__ import annotations

from docx import Document as DocxDocument

from app.knowledge.extractors.base import BaseExtractor, ExtractedContent


class DocxExtractor(BaseExtractor):
    HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"}

    def extract(self, file_path: str, filename: str) -> ExtractedContent:
        try:
            doc = DocxDocument(file_path)
        except Exception as exc:
            return ExtractedContent(
                text=f"[ERROR] Could not read DOCX: {exc}",
                metadata={"filename": filename, "error": str(exc)},
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

        paragraphs_text: list[str] = []
        sections: list[dict[str, object]] = []
        current_section = ""

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                paragraphs_text.append("")
                continue

            style_name = para.style.name if para.style else ""
            if style_name in self.HEADING_STYLES:
                current_section = text
                level = int(style_name.split()[-1]) if style_name[-1].isdigit() else 1
                sections.append(
                    {
                        "level": level,
                        "title": text,
                        "style": style_name,
                    }
                )

            paragraphs_text.append(text)

        full_text = "\n\n".join(paragraphs_text)

        return ExtractedContent(
            text=full_text,
            metadata={
                "filename": filename,
                "paragraph_count": len(doc.paragraphs),
                "section_count": len(sections),
                "sections": sections,
                "char_count": len(full_text),
            },
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    def supported_types(self) -> list[str]:
        return ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]
